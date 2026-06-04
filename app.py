from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
import openpyxl
from openpyxl import load_workbook
import os, io, base64, datetime, copy

app = Flask(__name__)
app.config['SERVER_NAME'] = None

CORS(app, origins='*', supports_credentials=False)

@app.before_request
def handle_host():
    pass  # allow all hosts

@app.after_request
def after_request(response):
    response.headers['Access-Control-Allow-Origin'] = '*'
    response.headers['Access-Control-Allow-Headers'] = 'Content-Type,Accept,Authorization'
    response.headers['Access-Control-Allow-Methods'] = 'GET,POST,PUT,OPTIONS'
    return response

@app.route('/', methods=['GET'])
def index():
    return jsonify({'status': 'ok', 'service': 'Southern Lakes Order API'})

OUR_NAME = 'Southern Lakes Shade & Screens'
OUR_ADDR = '37a Sargood Road, Wanaka'
OUR_PHONE = '021 224 0908'

def safe_set(ws, addr, val):
    try:
        cell = ws[addr]
        if hasattr(cell, 'value'):
            cell.value = val
    except:
        pass

def get_template(name):
    # Try templates subfolder first, then same directory
    base = os.path.dirname(os.path.abspath(__file__))
    path = os.path.join(base, 'templates', name)
    if not os.path.exists(path):
        path = os.path.join(base, name)
    return load_workbook(path)

@app.route('/health', methods=['GET'])
def health():
    return jsonify({'status': 'ok'})

@app.route('/order/honeycomb', methods=['POST'])
def honeycomb():
    data = request.json
    wb = get_template('honeycomb.xlsx')
    ws = wb['Sheet1']
    safe_set(ws, 'I1', OUR_NAME)
    safe_set(ws, 'I3', OUR_ADDR)
    safe_set(ws, 'I7', OUR_PHONE)
    safe_set(ws, 'Q5', data.get('custName',''))
    safe_set(ws, 'Q7', data.get('orderDate',''))
    blinds = data.get('items', [])
    rows = [13,14,15,16,17,18,19,20,21,22]
    for i, row in enumerate(rows):
        if i < len(blinds):
            b = blinds[i]
            safe_set(ws, f'A{row}', i+1)
            safe_set(ws, f'B{row}', b.get('location',''))
            safe_set(ws, f'C{row}', int(b['width']) if b.get('width') else '')
            safe_set(ws, f'E{row}', int(b['drop']) if b.get('drop') else '')
            safe_set(ws, f'G{row}', b.get('controlType','') + ' Drive')
            safe_set(ws, f'I{row}', b.get('fabricType','').replace('_',' '))
            safe_set(ws, f'K{row}', b.get('fabricColour',''))
            safe_set(ws, f'N{row}', b.get('railColour',''))
            safe_set(ws, f'O{row}', b.get('componentColour','') if b.get('controlType')=='Chain' else 'Clear')
            safe_set(ws, f'P{row}', b.get('chainLength','') if b.get('controlType')=='Chain' else '')
            safe_set(ws, f'Q{row}', b.get('controlSide',''))
            st = b.get('sideTracks','No')
            if st == 'Yes' and b.get('sideTrackDrop'):
                st += f' ({b["sideTrackDrop"]}mm)'
            safe_set(ws, f'R{row}', st)
    notes = data.get('notes','')
    if notes:
        safe_set(ws, 'A28', 'Notes: ' + notes)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"{data.get('custName','order')}_Honeycomb_Order.xlsx"
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/order/pleated', methods=['POST'])
def pleated():
    data = request.json
    wb = get_template('pleated.xlsx')
    ws = wb['Sheet1']
    safe_set(ws, 'I1', OUR_NAME)
    safe_set(ws, 'I3', OUR_ADDR)
    safe_set(ws, 'I7', OUR_PHONE)
    safe_set(ws, 'Q5', data.get('custName',''))
    safe_set(ws, 'Q7', data.get('orderDate',''))
    angle_labels = {'1225':'12x25','2020':'20x20','2030':'20x30','2040':'20x40','2050':'20x50','2570':'25x70','uchannel':'U Channel','custom':'Custom'}
    side_labels = {'all4':'all sides','3sides_top_bottom_left':'T+B+L','3sides_top_bottom_right':'T+B+R','top_bottom':'T+B','left_right':'L+R','top_only':'Top','bottom_only':'Bottom','left_only':'Left','right_only':'Right'}
    items = data.get('items', [])
    rows = [13,14,15,16,17,18,19,20,21,22]
    for i, row in enumerate(rows):
        if i < len(items):
            p = items[i]
            safe_set(ws, f'A{row}', p.get('location', i+1))
            safe_set(ws, f'C{row}', int(p['height']) if p.get('height') else '')
            safe_set(ws, f'E{row}', int(p['track']) if p.get('track') else '')
            safe_set(ws, f'G{row}', '' if p.get('sdbl')=='double' else 'X')
            safe_set(ws, f'I{row}', 'X' if p.get('sdbl')=='double' else '')
            safe_set(ws, f'K{row}', 'Internal' if p.get('fit')=='Internal' else 'Face Fit')
            asize = p.get('angleSize','none')
            if asize and asize != 'none':
                aside = side_labels.get(p.get('angleSides','all4'), '')
                safe_set(ws, f'M{row}', f"{angle_labels.get(asize,asize)} — {aside}")
            safe_set(ws, f'R{row}', p.get('trimColour',''))
    pcs = next((p for p in items if p.get('powderCoat')), None)
    if pcs:
        safe_set(ws, 'E28', pcs['powderCoat'])
    safe_set(ws, 'E29', 'Dylan')
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"{data.get('custName','order')}_Pleated_Order.xlsx"
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/order/rollaway', methods=['POST'])
def rollaway():
    data = request.json
    wb = get_template('rollaway.xlsx')
    ws = wb[wb.sheetnames[0]]
    safe_set(ws, 'C3', data.get('custName',''))
    safe_set(ws, 'I3', OUR_NAME)
    ws['C5'] = data.get('orderDate','')  # override TODAY() formula
    safe_set(ws, 'I7', OUR_PHONE)
    angle_labels = {'1225':'12x25','2020':'20x20','2030':'20x30','2040':'20x40','2050':'20x50','2570':'25x70','uchannel':'U Channel','custom':'Custom'}
    side_labels = {'all4':'all sides','3sides_top_bottom_left':'T+B+L','3sides_top_bottom_right':'T+B+R','top_bottom':'T+B','left_right':'L+R','top_only':'Top','bottom_only':'Bottom','left_only':'Left','right_only':'Right'}
    items = data.get('items', [])
    rows = [10,11,12,13,14,15,16,17,18,19,20,21,22,23,24]
    for i, row in enumerate(rows):
        if i < len(items):
            r = items[i]
            safe_set(ws, f'A{row}', r.get('location',''))
            safe_set(ws, f'B{row}', int(r['cassette']) if r.get('cassette') else '')
            safe_set(ws, f'D{row}', int(r['track']) if r.get('track') else '')
            safe_set(ws, f'F{row}', '' if r.get('topFix')=='NA' else r.get('topFix',''))
            cord = r.get('cord','')
            safe_set(ws, f'G{row}', 'B' if cord=='Black' else 'W' if cord=='White' else '')
            safe_set(ws, f'H{row}', 'Y' if r.get('foot')=='Small' else '')
            asize = r.get('angleSize','none')
            if asize and asize != 'none':
                aside = side_labels.get(r.get('angleSides','all4'), '')
                astr = f"{angle_labels.get(asize,asize)} {aside}"
                if asize == '2020': safe_set(ws, f'I{row}', astr)
                elif asize == '2030': safe_set(ws, f'K{row}', astr)
                elif asize == '2040': safe_set(ws, f'M{row}', astr)
                else: safe_set(ws, f'I{row}', astr)
            safe_set(ws, f'P{row}', r.get('colour',''))
    pcs = next((r for r in items if r.get('powderCoat') or r.get('customColourName')), None)
    if pcs:
        safe_set(ws, 'D30', pcs.get('powderCoat') or pcs.get('customColourName',''))
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"{data.get('custName','order')}_Rollaway_Order.xlsx"
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)
